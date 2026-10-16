#!/usr/bin/env python3
"""
Vigia do PIB dos Municípios (IBGE).

O IBGE publica o PIB municipal uma vez por ano, em dezembro, sempre com dois
anos de defasagem: a edição divulgada em dezembro de 2025 trouxe 2023. Em vez
de tentar adivinhar a data, este script olha a listagem do FTP e compara o ano
mais recente com o que já foi processado aqui.

A base vem num único .zip com a série inteira (2010 em diante), então não há
como baixar "só o ano novo" — mas são 20 MB, e o download só acontece quando
há ano novo de verdade.

    python scripts/monitor_pib.py --so-checar    # só diz se mudou
    python scripts/monitor_pib.py                # baixa e grava fonte/auto/
    python scripts/monitor_pib.py --forcar       # reprocessa mesmo sem ano novo

Saída: fonte/auto/auto_pib.csv + auto_pib.meta.json + auto_pib.estado.json
"""
import argparse, io, json, os, re, sys, zipfile
from datetime import datetime, timezone
import urllib.request

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
AUTO = os.path.join(RAIZ, 'fonte', 'auto')
FTP = 'https://ftp.ibge.gov.br/Pib_Municipios/'
PAGINA = ('https://www.ibge.gov.br/estatisticas/economicas/contas-nacionais/'
          '9088-produto-interno-bruto-dos-municipios.html')

# Como cada indicador do painel sai da planilha do IBGE. O rótulo da coluna na
# origem é longo e quebrado em várias linhas ("Valor adicionado bruto da\n
# Agropecuária, ..."), então o casamento é por trecho, não por nome exato: o
# IBGE reescreve esses cabeçalhos de uma edição para outra.
COLUNAS = {
    'vab_agro':  ('valor adicionado bruto da agropecuaria', 1000),
    'vab_ind':   ('valor adicionado bruto da industria', 1000),
    'vab_serv':  ('valor adicionado bruto dos servicos', 1000),
    'vab_adm':   ('valor adicionado bruto da administracao', 1000),
    'vab_total': ('valor adicionado bruto total', 1000),
    'pib_total': ('produto interno bruto,', 1000),
    'pib_pc':    ('produto interno bruto per capita', 1),
}
# Participações: calculadas aqui, não lidas da origem — o IBGE não as publica.
PARTICIPACOES = {
    'part_vab_agro': 'vab_agro',
    'part_vab_ind':  'vab_ind',
    'part_vab_serv': 'vab_serv',
    'part_vab_adm':  'vab_adm',
}
ORDEM = (list(COLUNAS) + list(PARTICIPACOES))

NOMES = {
    'pib_total': 'PIB total', 'pib_pc': 'PIB per capita',
    'vab_total': 'VAB total', 'vab_agro': 'VAB agropecuária',
    'vab_ind': 'VAB indústria', 'vab_adm': 'VAB administração pública',
    'vab_serv': 'VAB serviços',
    'part_vab_agro': 'Participação da agropecuária no VAB',
    'part_vab_ind': 'Participação da indústria no VAB',
    'part_vab_adm': 'Participação da adm. pública no VAB',
    'part_vab_serv': 'Participação dos serviços no VAB',
}


def chave(t):
    """Texto comparável: sem acento, sem quebra de linha, em minúsculas."""
    import unicodedata
    t = unicodedata.normalize('NFKD', str(t or '')).encode('ascii', 'ignore').decode()
    return ' '.join(t.lower().split())


def busca(url, binario=False):
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req, timeout=300) as r:
        dados = r.read()
    return dados if binario else dados.decode('utf-8', 'ignore')


def ultima_edicao():
    """
    Ano mais recente e endereço do .zip com a série municipal.

    As pastas do FTP mudam de nome conforme a edição: já foram '2021/' e
    '2022_2023/'. O que se procura é o maior ano que aparece em qualquer nome
    de pasta — assim o script sobrevive à próxima troca de padrão.
    """
    html = busca(FTP)
    pastas = [m for m in re.findall(r'href="([^"]+/)"', html)
              if re.fullmatch(r'\d{4}(_\d{4})?/', m)]
    if not pastas:
        raise SystemExit('Não consegui listar as pastas em ' + FTP +
                         '. O IBGE mudou a estrutura do FTP.')
    def ano_final(p):
        return max(int(a) for a in re.findall(r'\d{4}', p))
    pasta = max(pastas, key=ano_final)
    ano = ano_final(pasta)

    dentro = busca(FTP + pasta + 'base/')
    zips = [m for m in re.findall(r'href="([^"]+\.zip)"', dentro)
            if 'xlsx' in m.lower() and str(ano) in m]
    if not zips:
        raise SystemExit('A pasta %sbase/ não trouxe o .zip xlsx da série. '
                         'Confira %s antes de mexer no script.' % (FTP + pasta, PAGINA))
    return ano, FTP + pasta + 'base/' + zips[0]


def ler_planilha(bruto, ano):
    """
    Valores por município, cada indicador no ano mais recente em que existe.

    O IBGE não publica todas as colunas até o mesmo ano: na edição 2022-2023,
    o PIB e o PIB per capita vão até 2023, mas o VAB por setor para no ano
    2021. Pegar tudo do último ano deixaria o VAB vazio; pegar tudo de 2021
    jogaria o PIB dois anos para trás. Então cada indicador é lido do ano mais
    recente em que a coluna tem valor, e esse ano vai junto no rótulo.
    """
    from openpyxl import load_workbook
    z = zipfile.ZipFile(io.BytesIO(bruto))
    nome = [n for n in z.namelist() if n.lower().endswith('.xlsx')]
    if not nome:
        raise SystemExit('O .zip do IBGE não trouxe nenhum .xlsx dentro.')
    wb = load_workbook(io.BytesIO(z.read(nome[0])), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    cab = [chave(c) for c in next(it)]

    def acha(trecho):
        for i, c in enumerate(cab):
            if c.startswith(trecho):
                return i
        raise SystemExit("Não achei a coluna que começa com '%s' na planilha do "
                         "IBGE. O cabeçalho mudou; confira %s." % (trecho, PAGINA))

    i_ano = acha('ano')
    i_cod = acha('codigo do municipio')
    pos = {ind: acha(trecho) for ind, (trecho, _) in COLUNAS.items()}

    # Uma passada só: guarda, para cada indicador, o valor de cada ano.
    por_ind = {ind: {} for ind in pos}
    for r in it:
        if not r or r[i_ano] is None:
            continue
        try:
            a = int(r[i_ano])
        except (TypeError, ValueError):
            continue
        if a > ano:
            continue
        cod = str(r[i_cod]).split('.')[0].strip()
        if len(cod) != 7:
            continue
        for ind, j in pos.items():
            if r[j] not in (None, ''):
                por_ind[ind].setdefault(a, {})[cod] = r[j]

    linhas, anos = {}, {}
    for ind, por_ano in por_ind.items():
        if not por_ano:
            print('  aviso: a coluna de %s veio vazia em todos os anos' % ind)
            continue
        recente = max(por_ano)
        anos[ind] = recente
        for cod, valor in por_ano[recente].items():
            linhas.setdefault(cod, {})[ind] = valor
    return linhas, anos


def montar(linhas, anos):
    """Converte para a unidade do painel e calcula as participações."""
    for part, origem in PARTICIPACOES.items():
        if origem in anos and 'vab_total' in anos:
            anos[part] = min(anos[origem], anos['vab_total'])
    saida = {}
    for cod, campos in linhas.items():
        v = {}
        for ind, (_, fator) in COLUNAS.items():
            bruto = campos.get(ind)
            if bruto is None or bruto == '':
                continue
            try:
                v[ind] = round(float(bruto) * fator, 2)
            except (TypeError, ValueError):
                continue
        total = v.get('vab_total')
        if total:
            for part, origem in PARTICIPACOES.items():
                if origem in v:
                    v[part] = round(v[origem] / total * 100, 2)
        saida[cod] = v
    return saida


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--so-checar', action='store_true')
    ap.add_argument('--forcar', action='store_true')
    args = ap.parse_args()

    os.makedirs(AUTO, exist_ok=True)
    caminho_estado = os.path.join(AUTO, 'auto_pib.estado.json')
    estado = {}
    if os.path.exists(caminho_estado):
        estado = json.loads(open(caminho_estado, encoding='utf-8').read())

    ano, url_zip = ultima_edicao()
    anterior = estado.get('ano')
    print('Ano mais recente no IBGE: %s' % ano)
    print('Último processado aqui: %s' % (anterior or 'nenhum'))

    novidade = (anterior != ano)
    if args.so_checar:
        print('MUDOU' if novidade else 'sem novidade')
        saida = os.environ.get('GITHUB_OUTPUT')
        if saida:
            with open(saida, 'a', encoding='utf-8') as f:
                f.write('novidade=%s\n' % ('true' if novidade else 'false'))
                f.write('ano=%s\n' % ano)
        return
    if not novidade and not args.forcar:
        print('Nada a fazer.')
        return

    print('Baixando %s' % url_zip.rsplit('/', 1)[-1])
    bruto = busca(url_zip, binario=True)
    print('  %.1f MB' % (len(bruto) / 1e6))
    linhas, anos = ler_planilha(bruto, ano)
    dados = montar(linhas, anos)
    if not dados:
        raise SystemExit('A planilha não trouxe nenhuma linha até o ano %s.' % ano)
    for ind in ORDEM:
        if ind in anos:
            print('  %-15s ano %s' % (ind, anos[ind]))

    # Conferência: o PIB do Brasil é a soma dos municípios. Um erro de unidade
    # (R$ mil lido como reais) apareceria aqui como três ordens de grandeza.
    soma = sum(v.get('pib_total', 0) for v in dados.values())
    print('  %d municípios · PIB somado: R$ %.2f tri' % (len(dados), soma / 1e12))
    if not 5e12 < soma < 2e13:
        print('  AVISO: o PIB somado saiu fora da faixa esperada para um ano '
              'recente (R$ 5 a 20 tri). Confira a unidade antes de publicar.')

    caminho_csv = os.path.join(AUTO, 'auto_pib.csv')
    with open(caminho_csv, 'w', encoding='utf-8', newline='') as f:
        f.write('codigo;' + ';'.join(ORDEM) + '\n')
        for cod in sorted(dados):
            v = dados[cod]
            f.write(cod + ';' + ';'.join(
                ('' if v.get(i) is None else repr(v[i])) for i in ORDEM) + '\n')

    # Cada indicador leva o ano em que de fato foi publicado — não o ano da
    # edição. É o que evita um VAB de 2021 aparecer no painel como se fosse
    # de 2023, que foi exatamente o engano que esta rotina veio corrigir.
    meta = {ind: {'nome': NOMES[ind], 'ano': str(anos[ind]),
                  'fonte': 'IBGE — PIB dos Municípios %s' % anos[ind]}
            for ind in ORDEM if ind in anos}
    with open(os.path.join(AUTO, 'auto_pib.meta.json'), 'w', encoding='utf-8') as f:
        json.dump(meta, f, ensure_ascii=False, indent=2)

    estado = {'ano': ano, 'anosPorIndicador': {i: anos[i] for i in ORDEM if i in anos},
              'arquivo': url_zip.rsplit('/', 1)[-1], 'origem': PAGINA,
              'verificadoEm': datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC'),
              'totais': {'municipios': len(dados), 'pib_somado': round(soma, 2)}}
    with open(caminho_estado, 'w', encoding='utf-8') as f:
        json.dump(estado, f, ensure_ascii=False, indent=2)
    print('Gravado em fonte/auto/auto_pib.csv (%d municípios)' % len(dados))


if __name__ == '__main__':
    main()
